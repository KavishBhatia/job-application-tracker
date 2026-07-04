import io

from openpyxl import Workbook, load_workbook

COLUMNS = ["id", "company", "role", "date_applied", "status", "job_post_url", "source_text"]


def to_excel(rows: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append([row.get(column, "") for column in COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def from_excel(file) -> list[dict]:
    content = file.read() if hasattr(file, "read") else file
    workbook = load_workbook(filename=io.BytesIO(content))
    sheet = workbook.active

    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter)
    return [
        {header[i]: row[i] for i in range(len(header)) if row[i] is not None}
        for row in rows_iter
    ]
